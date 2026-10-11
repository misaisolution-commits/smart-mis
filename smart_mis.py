"""
Smart MIS - upload data, type a prompt, get a report.

How it works
1. Your file is cleaned (numbers/dates auto-detected).
2. The AI (Ollama) writes pandas code for your question and the code is
   run on the FULL dataset, so totals/averages are exact (auto-retries on error).
3. The AI then writes a management report using only those computed numbers
   plus an automatic data fact-sheet.
4. Download the report as Word (.docx), PDF, HTML, Markdown or Excel.

Run locally:  streamlit run smart_mis.py
Deploy:       see DEPLOY.md (needs requirements.txt)

Optional Streamlit secrets (.streamlit/secrets.toml or the Cloud "Secrets" box):
    APP_PASSWORD = "your-password"          # turns on a password gate
    OLLAMA_URL        = "https://your-tunnel-url"  # default Ollama URL
    ANTHROPIC_API_KEY = "sk-ant-..."               # use Claude API (no tunnel)
    ANTHROPIC_MODEL   = "claude-sonnet-5-5"        # optional
"""

import builtins
import hmac
import html
import io
import json
import re
import time
from datetime import datetime

import numpy as np
import pandas as pd
import plotly.express as px
import requests
import streamlit as st

st.set_page_config(page_title="Smart MIS", page_icon="📊", layout="wide")


def _secret(name: str, default=None):
    """Read a Streamlit secret without crashing when no secrets file exists."""
    try:
        return st.secrets[name]
    except Exception:  # noqa: BLE001
        return default


def password_gate() -> bool:
    """If APP_PASSWORD is set in secrets, require it before showing the app."""
    expected = _secret("APP_PASSWORD")
    if not expected:
        return True
    if st.session_state.get("auth_ok"):
        return True
    st.title("🔒 Smart MIS")
    pw = st.text_input("Password", type="password")
    if pw:
        if hmac.compare_digest(pw.encode(), str(expected).encode()):
            st.session_state["auth_ok"] = True
            st.rerun()
        else:
            st.error("Incorrect password.")
    return False


if not password_gate():
    st.stop()

# --------------------------------------------------------------------------
# Data loading & cleaning
# --------------------------------------------------------------------------


@st.cache_resource(show_spinner=False)  # resource cache = no copy of big data on every click
def load_file(name: str, data: bytes) -> dict:
    bio = io.BytesIO(data)
    if name.lower().endswith(".csv"):
        try:
            return {"Sheet1": pd.read_csv(bio)}
        except UnicodeDecodeError:
            bio.seek(0)
            return {"Sheet1": pd.read_csv(bio, encoding="latin-1")}
    return pd.read_excel(bio, sheet_name=None)


def _to_num(s: pd.Series) -> pd.Series:
    return pd.to_numeric(
        s.str.replace(",", "", regex=False)
        .str.replace("₹", "", regex=False)
        .str.replace("$", "", regex=False),
        errors="coerce",
    )


def clean(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    df = df.dropna(how="all").dropna(axis=1, how="all")
    for c in df.columns:
        # pandas 3 stores text as "str" (not "object"), so accept both
        if not (pd.api.types.is_object_dtype(df[c]) or pd.api.types.is_string_dtype(df[c])):
            continue
        nonnull = df[c].notna().sum()
        if nonnull == 0:
            continue
        # Speed: test a 200-value sample first, so text columns (names, remarks...)
        # skip the expensive full-column conversion attempts.
        samp = df[c].dropna().head(200).astype(str).str.strip()
        if _to_num(samp).notna().mean() > 0.9:
            num = _to_num(df[c].astype(str).str.strip())
            if num.notna().sum() / nonnull > 0.9:
                df[c] = num
                continue
        if re.search(r"date|time|month|period|day", c, re.I):
            # Detect ISO dates first to avoid dayfirst warnings for YYYY-MM-DD.
            iso_ratio = samp.str.match(r"^\d{4}-\d{1,2}-\d{1,2}(?:$|[ T])").mean()
            is_iso = iso_ratio > 0.8
            sample_dates = pd.to_datetime(
                samp, errors="coerce", dayfirst=not is_iso, yearfirst=is_iso
            )
            if sample_dates.notna().mean() > 0.8:
                dt = pd.to_datetime(
                    df[c].astype(str).str.strip(),
                    errors="coerce",
                    dayfirst=not is_iso,
                    yearfirst=is_iso,
                )
                if dt.notna().sum() / nonnull > 0.8:
                    df[c] = dt
    return df.reset_index(drop=True)


def column_types(df: pd.DataFrame):
    num = df.select_dtypes(include="number").columns.tolist()
    dates = df.select_dtypes(include=["datetime", "datetimetz"]).columns.tolist()
    cats = [
        c
        for c in df.columns
        if c not in num + dates and 1 < df[c].nunique() <= 100
    ]
    return num, dates, cats


# --------------------------------------------------------------------------
# Fact sheet (always sent to the AI)
# --------------------------------------------------------------------------


def _d(series: pd.Series) -> dict:
    return {str(k): round(float(v), 2) for k, v in series.items()}


def build_facts(df: pd.DataFrame) -> dict:
    num, dates, cats = column_types(df)
    facts = {
        "rows": len(df),
        "columns": {c: str(t) for c, t in df.dtypes.items()},
        "missing_values": {c: int(v) for c, v in df.isna().sum().items() if v},
    }
    if num:
        desc = df[num].describe().T
        facts["numeric_summary"] = {
            c: {
                "total": round(float(df[c].sum()), 2),
                "mean": round(float(desc.loc[c, "mean"]), 2),
                "min": round(float(desc.loc[c, "min"]), 2),
                "median": round(float(desc.loc[c, "50%"]), 2),
                "max": round(float(desc.loc[c, "max"]), 2),
            }
            for c in num[:10]
        }
    breakdown = {}
    for cat in cats[:5]:
        for n in num[:3]:
            g = df.groupby(cat)[n].sum().sort_values(ascending=False)
            breakdown[f"{n} by {cat}"] = {
                "top": _d(g.head(7)),
                "bottom": _d(g.tail(3)),
            }
    if breakdown:
        facts["breakdowns"] = breakdown
    if dates and num:
        d = dates[0]
        m = (
            df.dropna(subset=[d])
            .set_index(d)[num[:3]]
            .resample("MS")
            .sum()
            .round(2)
            .tail(24)
        )
        m.index = m.index.strftime("%Y-%m")
        facts[f"monthly_trend_by_{d}"] = m.to_dict("index")
    facts["duplicate_rows"] = int(df.duplicated().sum())
    return facts


# --------------------------------------------------------------------------
# AI (Ollama)
# --------------------------------------------------------------------------


@st.cache_data(ttl=15, show_spinner=False)  # don't ping Ollama on every click
def ollama_models(url: str) -> list:
    try:
        r = requests.get(f"{url}/api/tags", timeout=5)
        r.raise_for_status()
        return [m["name"] for m in r.json().get("models", [])]
    except Exception:  # noqa: BLE001
        return []


KEEP_ALIVE = "30m"  # keep the model loaded in memory between requests
NUM_CTX = 8192  # keep constant: changing it makes Ollama reload the model


def _payload(model, prompt, temperature, max_tokens, stream) -> dict:
    return {
        "model": model,
        "prompt": prompt,
        "stream": stream,
        "keep_alive": KEEP_ALIVE,
        "options": {
            "temperature": temperature,
            "num_ctx": NUM_CTX,
            "num_predict": max_tokens,  # hard cap so the model can't ramble
        },
    }


def _ollama_llm(url: str, model: str, prompt: str, temperature: float = 0.2,
                max_tokens: int = 600) -> str:
    r = requests.post(
        f"{url}/api/generate",
        json=_payload(model, prompt, temperature, max_tokens, False),
        timeout=600,
    )
    r.raise_for_status()
    return r.json()["response"]


def _ollama_stream(url: str, model: str, prompt: str, temperature: float = 0.3,
                   max_tokens: int = 900):
    """Yield the answer piece by piece so the report appears while it is written."""
    with requests.post(
        f"{url}/api/generate",
        json=_payload(model, prompt, temperature, max_tokens, True),
        stream=True,
        timeout=600,
    ) as r:
        r.raise_for_status()
        for line in r.iter_lines():
            if not line:
                continue
            chunk = json.loads(line)
            if chunk.get("error"):
                raise RuntimeError(chunk["error"])
            if chunk.get("response"):
                yield chunk["response"]
            if chunk.get("done"):
                break


# ---- Claude API (online, no tunnel needed) --------------------------------

ANTHROPIC_URL = "https://api.anthropic.com/v1/messages"
PROVIDER = {"name": "ollama"}  # set from the sidebar on every run


def _claude_headers() -> dict:
    key = _secret("ANTHROPIC_API_KEY")
    if not key:
        raise RuntimeError(
            "ANTHROPIC_API_KEY is missing. Add it in Streamlit Settings > Secrets."
        )
    return {
        "x-api-key": str(key),
        "anthropic-version": "2023-06-01",
        "content-type": "application/json",
    }


def _claude_body(model, prompt, temperature, max_tokens, stream) -> dict:
    return {
        "model": model,
        "max_tokens": max(max_tokens, 1024),
        "stream": stream,
        "messages": [{"role": "user", "content": prompt}],
    }


def _claude_llm(model: str, prompt: str, temperature: float, max_tokens: int) -> str:
    r = requests.post(
        ANTHROPIC_URL,
        headers=_claude_headers(),
        json=_claude_body(model, prompt, temperature, max_tokens, False),
        timeout=300,
    )
    if r.status_code >= 400:
        raise RuntimeError(f"Claude API error {r.status_code}: {r.text[:300]}")
    return "".join(b.get("text", "") for b in r.json().get("content", []))


def _claude_stream(model: str, prompt: str, temperature: float, max_tokens: int):
    with requests.post(
        ANTHROPIC_URL,
        headers=_claude_headers(),
        json=_claude_body(model, prompt, temperature, max_tokens, True),
        stream=True,
        timeout=300,
    ) as r:
        if r.status_code >= 400:
            raise RuntimeError(f"Claude API error {r.status_code}: {r.text[:300]}")
        for raw in r.iter_lines():
            if not raw:
                continue
            line = raw.decode("utf-8") if isinstance(raw, bytes) else raw
            if not line.startswith("data:"):
                continue
            data = line[5:].strip()
            if not data or data == "[DONE]":
                continue
            ev = json.loads(data)
            if ev.get("type") == "error":
                raise RuntimeError(ev.get("error", {}).get("message", "Claude API error"))
            if ev.get("type") == "content_block_delta":
                txt = ev.get("delta", {}).get("text")
                if txt:
                    yield txt


# ---- Dispatchers used by the rest of the app ------------------------------


def llm(url: str, model: str, prompt: str, temperature: float = 0.2,
        max_tokens: int = 600) -> str:
    if PROVIDER["name"] == "claude":
        return _claude_llm(model, prompt, temperature, max_tokens)
    return _ollama_llm(url, model, prompt, temperature, max_tokens)


def llm_stream(url: str, model: str, prompt: str, temperature: float = 0.3,
               max_tokens: int = 900):
    if PROVIDER["name"] == "claude":
        yield from _claude_stream(model, prompt, temperature, max_tokens)
    else:
        yield from _ollama_stream(url, model, prompt, temperature, max_tokens)


# --------------------------------------------------------------------------
# AI-written pandas code, executed on the full data
# NOTE: the blocklist below is a basic guard, NOT a real sandbox. Keep the
# deployed app private (password / allow-listed viewers).
# --------------------------------------------------------------------------

_ALLOWED = [
    "len", "sum", "min", "max", "abs", "round", "sorted", "range", "list",
    "dict", "set", "tuple", "str", "int", "float", "bool", "enumerate",
    "zip", "map", "filter", "any", "all", "isinstance", "reversed", "pow",
    "divmod",
]
SAFE_BUILTINS = {k: getattr(builtins, k) for k in _ALLOWED}

BLOCKED = re.compile(
    r"__|\bimport\b|\bopen\s*\(|\bexec\b|\beval\b|\bos\.|\bsys\.|subprocess|"
    r"getattr|setattr|globals|locals|\binput\s*\(|compile|\bwhile\b|"
    r"to_csv|to_excel|to_pickle|to_sql|to_parquet|to_json|to_clipboard|"
    r"to_hdf|to_feather|read_|np\.(load|save|fromfile|memmap|genfromtxt|"
    r"loadtxt|savetxt)|pd\.(io|compat|util)\b"
)


def extract_code(text: str) -> str:
    m = re.search(r"```(?:python)?\s*\n(.*?)```", text, re.S)
    return (m.group(1) if m else text).strip()


def run_code(code: str, df: pd.DataFrame):
    if BLOCKED.search(code):
        raise ValueError("Code contains blocked operations")
    env = {"__builtins__": SAFE_BUILTINS, "df": df.copy(), "pd": pd, "np": np}
    exec(code, env)  # noqa: S102
    if "result" not in env:
        raise ValueError("Code did not define a variable named `result`")
    return env["result"]


def result_to_text(r) -> str:
    if isinstance(r, (pd.DataFrame, pd.Series)):
        return r.head(60).to_string()
    return str(r)[:4000]


def compute_answer(url, model, df, question):
    schema = "\n".join(f"- {c} ({t})" for c, t in df.dtypes.astype(str).items())
    sample = df.head(3).astype(str).to_string(index=False)
    base = f"""You write pandas code. A DataFrame named `df` is already loaded; `pd` and `np` are available.

Columns:
{schema}

Sample rows:
{sample}

Write Python code that computes the exact numbers needed for this request:
"{question}"

Rules:
- Put the final answer in a variable named `result` (DataFrame, Series or number).
- No imports, no file access, no print statements, no while loops.
- Use the exact column names above.
- Reply with ONLY one ```python code block."""
    err, code = None, ""
    for _ in range(3):
        prompt = base
        if err:
            prompt += (
                f"\n\nYour previous code:\n```python\n{code}\n```\n"
                f"failed with: {err}\nFix it."
            )
        code = extract_code(llm(url, model, prompt, 0.1, max_tokens=400))
        try:
            return code, result_to_text(run_code(code, df)), None
        except Exception as e:  # noqa: BLE001
            err = f"{type(e).__name__}: {e}"
    return code, None, err


def slim_facts(facts: dict) -> dict:
    """Short fact sheet (fewer prompt tokens = faster) used when exact numbers exist."""
    keep = ("rows", "columns", "missing_values", "numeric_summary", "duplicate_rows")
    return {k: v for k, v in facts.items() if k in keep}


def write_report_stream(url, model, question, computed, facts):
    if computed:
        facts = slim_facts(facts)
    prompt = f"""You are a senior business MIS analyst. Write a clear management report in Markdown.

REQUEST: {question}

Rules:
- Use ONLY the numbers given below. Never invent figures.
- If something cannot be determined from the data, say so.
- Be specific: quote numbers, names and percentages.

Structure:
## Summary  (2-3 sentences)
## Key Findings  (bullets with numbers)
## Risks / Anomalies
## Recommendations

COMPUTED RESULT (exact, from the full dataset):
{computed or "Not available - rely on the fact sheet only."}

DATA FACT SHEET:
{json.dumps(facts, default=str, ensure_ascii=False)}
"""
    yield from llm_stream(url, model, prompt, 0.3)


# --------------------------------------------------------------------------
# Report export helpers (Markdown / HTML / Word / PDF / Excel)
# --------------------------------------------------------------------------

_BOLD = re.compile(r"\*\*(.+?)\*\*")


def parse_md(md: str) -> list:
    """Very small Markdown parser -> list of blocks.

    Block kinds: ("heading", level, text), ("bullet", text),
    ("para", text), ("table", rows)
    """
    blocks, table = [], []

    def flush_table():
        nonlocal table
        if table:
            blocks.append(("table", table))
            table = []

    for raw in md.splitlines():
        line = raw.strip()
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue  # separator row
            table.append(cells)
            continue
        flush_table()
        if not line or re.fullmatch(r"[-*_]{3,}", line):
            continue
        h = re.match(r"(#{1,4})\s+(.*)", line)
        if h:
            blocks.append(("heading", len(h.group(1)), h.group(2).strip("* ")))
            continue
        b = re.match(r"(?:[-*•]|\d+[.)])\s+(.*)", line)
        if b:
            blocks.append(("bullet", b.group(1)))
            continue
        blocks.append(("para", line))
    flush_table()
    return blocks


def _plain(text: str) -> str:
    return text.replace("`", "")


def report_md_file(title: str, question: str, md: str, computed) -> bytes:
    out = f"# {title}\n\n*Generated {datetime.now():%d %b %Y, %H:%M}*\n\n"
    out += f"**Request:** {question}\n\n---\n\n{md}\n"
    if computed:
        out += f"\n---\n\n## Appendix: Verified numbers\n\n```\n{computed}\n```\n"
    return out.encode("utf-8")


def report_html(title: str, question: str, md: str, computed) -> bytes:
    def inline(t):
        t = html.escape(_plain(t))
        return _BOLD.sub(r"<b>\1</b>", t)

    parts, in_list = [], False
    for blk in parse_md(md):
        if blk[0] != "bullet" and in_list:
            parts.append("</ul>")
            in_list = False
        if blk[0] == "heading":
            lvl = min(blk[1] + 1, 4)
            parts.append(f"<h{lvl}>{inline(blk[2])}</h{lvl}>")
        elif blk[0] == "bullet":
            if not in_list:
                parts.append("<ul>")
                in_list = True
            parts.append(f"<li>{inline(blk[1])}</li>")
        elif blk[0] == "table":
            rows = blk[1]
            t = "<table><thead><tr>" + "".join(f"<th>{inline(c)}</th>" for c in rows[0]) + "</tr></thead><tbody>"
            for r in rows[1:]:
                t += "<tr>" + "".join(f"<td>{inline(c)}</td>" for c in r) + "</tr>"
            parts.append(t + "</tbody></table>")
        else:
            parts.append(f"<p>{inline(blk[1])}</p>")
    if in_list:
        parts.append("</ul>")
    if computed:
        parts.append("<h2>Appendix: Verified numbers</h2>")
        parts.append(f"<pre>{html.escape(computed)}</pre>")

    css = (
        "body{font-family:'Segoe UI',Arial,sans-serif;max-width:900px;margin:30px auto;"
        "padding:0 20px;color:#222;line-height:1.55}"
        "h1{color:#1f3a5f;border-bottom:3px solid #1f3a5f;padding-bottom:8px}"
        "h2,h3{color:#1f3a5f;margin-top:26px}"
        ".meta{color:#666;font-size:13px}"
        ".req{background:#eef3f9;padding:10px 14px;border-left:4px solid #1f3a5f;margin:14px 0}"
        "table{border-collapse:collapse;width:100%;margin:12px 0}"
        "th,td{border:1px solid #ccc;padding:6px 10px;text-align:left}"
        "th{background:#1f3a5f;color:#fff}"
        "pre{background:#f5f5f5;padding:12px;overflow-x:auto;font-size:12px}"
        "@media print{body{margin:0}}"
    )
    doc = (
        f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>{html.escape(title)}</title>"
        f"<style>{css}</style></head><body>"
        f"<h1>{html.escape(title)}</h1>"
        f"<div class='meta'>Generated {datetime.now():%d %b %Y, %H:%M}</div>"
        f"<div class='req'><b>Request:</b> {html.escape(question)}</div>"
        + "".join(parts)
        + "</body></html>"
    )
    return doc.encode("utf-8")


def report_docx(title: str, question: str, md: str, computed):
    """Word document. Returns bytes, or None if python-docx is missing."""
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        return None

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    def add_runs(par, text):
        # FIXED: single-escaped regex so **bold** is detected correctly
        for chunk in re.split(r"(\*\*.+?\*\*)", _plain(text)):
            if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
                par.add_run(chunk[2:-2]).bold = True
            elif chunk:
                par.add_run(chunk)

    doc.add_heading(title, 0)
    meta = doc.add_paragraph(f"Generated {datetime.now():%d %b %Y, %H:%M}")
    meta.runs[0].italic = True
    p = doc.add_paragraph()
    p.add_run("Request: ").bold = True
    p.add_run(question)

    for blk in parse_md(md):
        if blk[0] == "heading":
            doc.add_heading(_plain(blk[2]), min(blk[1], 3))
        elif blk[0] == "bullet":
            add_runs(doc.add_paragraph(style="List Bullet"), blk[1])
        elif blk[0] == "table":
            rows = blk[1]
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol)
            t.style = "Light Grid Accent 1"
            for i, r in enumerate(rows):
                for j in range(ncol):
                    t.cell(i, j).text = _plain(r[j]) if j < len(r) else ""
        else:
            add_runs(doc.add_paragraph(), blk[1])

    if computed:
        doc.add_heading("Appendix: Verified numbers", 1)
        para = doc.add_paragraph()
        run = para.add_run(computed)
        run.font.name = "Consolas"
        run.font.size = Pt(8)

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def report_pdf(title: str, question: str, md: str, computed):
    """PDF document. Returns bytes, or None if reportlab is missing."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            ListFlowable, ListItem, Paragraph, Preformatted,
            SimpleDocTemplate, Spacer, Table, TableStyle,
        )
    except ImportError:
        return None

    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["BodyText"], fontSize=10, leading=14)
    h1 = ParagraphStyle("h1", parent=ss["Title"], textColor=colors.HexColor("#1f3a5f"), alignment=0)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], textColor=colors.HexColor("#1f3a5f"), spaceBefore=12)
    h3 = ParagraphStyle("h3", parent=ss["Heading3"], textColor=colors.HexColor("#1f3a5f"))
    mono = ParagraphStyle("mono", parent=ss["Code"], fontSize=7, leading=9)

    def inline(t):
        return _BOLD.sub(r"<b>\1</b>", html.escape(_plain(t)))

    story = [
        Paragraph(html.escape(title), h1),
        Paragraph(f"<i>Generated {datetime.now():%d %b %Y, %H:%M}</i>", body),
        Spacer(1, 4),
        Paragraph(f"<b>Request:</b> {html.escape(question)}", body),
        Spacer(1, 6),
    ]

    bullets = []

    def flush_bullets():
        nonlocal bullets
        if bullets:
            story.append(
                ListFlowable(
                    [ListItem(Paragraph(inline(b), body), leftIndent=12) for b in bullets],
                    bulletType="bullet", start="•", leftIndent=14,
                )
            )
            bullets = []

    for blk in parse_md(md):
        if blk[0] == "bullet":
            bullets.append(blk[1])
            continue
        flush_bullets()
        if blk[0] == "heading":
            story.append(Paragraph(inline(blk[2]), h2 if blk[1] <= 2 else h3))
        elif blk[0] == "table":
            rows = blk[1]
            ncol = max(len(r) for r in rows)
            data = [
                [Paragraph(inline(r[j]) if j < len(r) else "", body) for j in range(ncol)]
                for r in rows
            ]
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dfe8f3")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story += [t, Spacer(1, 6)]
        else:
            story.append(Paragraph(inline(blk[1]), body))
    flush_bullets()

    if computed:
        story.append(Paragraph("Appendix: Verified numbers", h2))
        story.append(Preformatted(computed, mono))

    bio = io.BytesIO()
    SimpleDocTemplate(
        bio, pagesize=A4, title=title,
        leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
    ).build(story)
    return bio.getvalue()


@st.cache_resource(show_spinner="Reading and cleaning your file...")
def prepare(name: str, data: bytes, sheet: str):
    """Clean the sheet + build the fact sheet ONCE per file (not on every click)."""
    df = clean(load_file(name, data)[sheet])
    return df, build_facts(df), column_types(df)


@st.cache_data(show_spinner=False, max_entries=10)
def make_exports(title: str, question: str, md: str, appendix) -> dict:
    """Build the report files once per report instead of on every click."""
    return {
        "pdf": report_pdf(title, question, md, appendix),
        "docx": report_docx(title, question, md, appendix),
        "html": report_html(title, question, md, appendix),
        "md": report_md_file(title, question, md, appendix),
    }


def report_excel(md: str, computed: str, df: pd.DataFrame) -> bytes:
    bio = io.BytesIO()
    out = df.copy()
    # Excel cannot store timezone-aware datetimes
    for c in out.select_dtypes(include=["datetimetz"]).columns:
        out[c] = out[c].dt.tz_localize(None)
    with pd.ExcelWriter(bio, engine="openpyxl") as xw:
        pd.DataFrame({"AI Report": md.split("\n")}).to_excel(
            xw, sheet_name="Report", index=False
        )
        if computed:
            pd.DataFrame({"Computed result": computed.split("\n")}).to_excel(
                xw, sheet_name="Computed", index=False
            )
        out.to_excel(xw, sheet_name="Data", index=False)
    return bio.getvalue()


# --------------------------------------------------------------------------
# UI
# --------------------------------------------------------------------------

st.title("📊 Smart MIS Reporting")
st.caption("Upload your data, write a prompt, get a report.")

with st.sidebar:
    st.header("⚙️ AI settings")
    has_key = bool(_secret("ANTHROPIC_API_KEY"))
    provider = st.radio(
        "AI provider",
        ["Ollama (free - local / tunnel)", "Claude API (paid - online)"],
        index=1 if has_key else 0,
        help="Ollama is free. Claude API is optional and paid (needs a key in Secrets).",
    )
    ollama_url = ""
    if provider.startswith("Claude"):
        PROVIDER["name"] = "claude"
        model = st.text_input(
            "Claude model", _secret("ANTHROPIC_MODEL", "claude-sonnet-5-5")
        ).strip()
        if has_key:
            st.success("Claude API key found.")
        else:
            st.warning("Add ANTHROPIC_API_KEY in Streamlit Settings > Secrets.")
    else:
        PROVIDER["name"] = "ollama"
        ollama_url = st.text_input(
            "Ollama URL", _secret("OLLAMA_URL", "http://localhost:11434")
        ).strip().rstrip("/")
        models = ollama_models(ollama_url)
        if models:
            model = st.selectbox("Model", models)
        else:
            model = st.text_input("Model", "qwen2.5:3b")
            st.warning(
                "Ollama not detected. Start it (or your tunnel), then refresh. "
                "On a hosted deployment, 'localhost' will NOT work - use your "
                "tunnel/server URL."
            )
    use_compute = st.toggle(
        "Calculate exact numbers (recommended)",
        value=True,
        help="AI writes pandas code and runs it on your full data.",
    )
    st.caption("Tip: `qwen2.5:7b` or `qwen2.5-coder:7b` gives much better results than 3b.")
    st.header("📥 Download options")
    include_appendix = st.checkbox(
        "Include verified numbers appendix", value=True,
        help="Adds the exact computed result to the downloaded report.",
    )
    report_title = st.text_input("Report title", "MIS Report")

uploaded = st.file_uploader("Upload Excel or CSV", type=["xlsx", "xls", "csv"])

if not uploaded:
    st.info("Upload a business Excel or CSV file to begin.")
    st.stop()

try:
    sheets = load_file(uploaded.name, uploaded.getvalue())
except Exception as e:  # noqa: BLE001
    st.error(f"Could not read the file: {e}")
    st.stop()

sheet = st.selectbox("Sheet", list(sheets)) if len(sheets) > 1 else list(sheets)[0]
df, facts, (num_cols, date_cols, cat_cols) = prepare(
    uploaded.name, uploaded.getvalue(), sheet
)

m1, m2, m3, m4 = st.columns(4)
m1.metric("Rows", f"{len(df):,}")
m2.metric("Columns", len(df.columns))
m3.metric("Numeric", len(num_cols))
m4.metric("Missing cells", f"{int(df.isna().sum().sum()):,}")

tab_report, tab_charts, tab_data = st.tabs(["🤖 AI Report", "📈 Charts", "🗂 Data"])

# ---------------- AI report ----------------
with tab_report:
    templates = {
        "Custom (type below)": "",
        "Full summary & trends": "Give a full management summary of this data with key trends and anomalies.",
        "Top & bottom performers": "Which categories perform best and worst by the main value columns? Show totals and share of total.",
        "Month-wise performance": "Show month-wise performance, growth versus previous month, and highlight the best and worst months.",
        "Data quality check": "Check data quality: missing values, duplicates, outliers and suspicious entries.",
    }
    choice = st.selectbox("Quick prompts", list(templates))
    question = st.text_area(
        "Your prompt",
        value=templates[choice],
        height=100,
        placeholder="e.g. Branch-wise sales total for last quarter and which branch is lagging?",
    )

    if st.button("Generate Report", type="primary"):
        if not question.strip():
            st.warning("Enter a prompt first.")
        else:
            try:
                computed, code, err, t_calc = None, None, None, 0.0
                if use_compute:
                    t0 = time.perf_counter()
                    with st.spinner("Calculating exact numbers from your data..."):
                        code, computed, err = compute_answer(
                            ollama_url, model, df, question
                        )
                    t_calc = time.perf_counter() - t0
                # Stream the report: text appears as soon as the model starts writing
                t0 = time.perf_counter()
                box = st.empty()
                with box.container():
                    st.markdown("### ✍️ Writing report...")
                    md = st.write_stream(
                        write_report_stream(ollama_url, model, question, computed, facts)
                    )
                box.empty()
                st.session_state["last"] = {
                    "md": md, "computed": computed, "code": code,
                    "err": err, "q": question,
                    "t_calc": t_calc, "t_rep": time.perf_counter() - t0,
                }
            except requests.RequestException as e:
                st.error(
                    "Could not reach the AI service. If using Ollama, make sure it "
                    "and your tunnel are running and the URL is set; if using the "
                    f"Claude API, check your internet/key. Details: {e}"
                )
            except RuntimeError as e:
                st.error(f"The AI returned an error: {e}")

    last = st.session_state.get("last")
    if last:
        st.markdown("### AI Report")
        st.markdown(last["md"])
        st.caption(
            f"⏱ Exact numbers: {last.get('t_calc', 0):.1f}s · "
            f"Report writing: {last.get('t_rep', 0):.1f}s"
        )

        if last["err"]:
            st.warning(
                "Exact calculation failed, so the report used the data fact sheet "
                f"only. Last error: {last['err']}"
            )
        if last["computed"]:
            with st.expander("🔍 Verify: exact numbers the report is based on"):
                st.code(last["computed"])
                st.caption("Code the AI ran on your full data:")
                st.code(last["code"], language="python")

        # ---------------- Download report ----------------
        st.markdown("### 📥 Download report")
        appendix = last["computed"] if include_appendix else None
        stamp = datetime.now().strftime("%Y%m%d_%H%M")
        base_name = re.sub(r"[^\w\-]+", "_", report_title).strip("_") or "MIS_Report"
        base_name = f"{base_name}_{stamp}"

        files = make_exports(report_title, last["q"], last["md"], appendix)
        docx_bytes, pdf_bytes = files["docx"], files["pdf"]

        d1, d2, d3, d4, d5 = st.columns(5)
        if pdf_bytes:
            d1.download_button(
                "⬇ PDF", pdf_bytes, f"{base_name}.pdf", "application/pdf",
                use_container_width=True,
            )
        else:
            d1.button("⬇ PDF", disabled=True, use_container_width=True,
                      help="Run: pip install reportlab")
        if docx_bytes:
            d2.download_button(
                "⬇ Word (.docx)", docx_bytes, f"{base_name}.docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
            )
        else:
            d2.button("⬇ Word (.docx)", disabled=True, use_container_width=True,
                      help="Run: pip install python-docx")
        d3.download_button(
            "⬇ HTML", files["html"],
            f"{base_name}.html", "text/html", use_container_width=True,
        )
        d4.download_button(
            "⬇ Markdown", files["md"],
            f"{base_name}.md", "text/markdown", use_container_width=True,
        )
        # Excel includes the whole dataset (slow to write), so build it only on demand
        if last.get("xlsx"):
            d5.download_button(
                "⬇ Excel + data", last["xlsx"], f"{base_name}.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
        elif d5.button(
            "📄 Prepare Excel + data", use_container_width=True,
            help="Builds the Excel file with your full data (slow for big files).",
        ):
            with st.spinner("Building Excel..."):
                last["xlsx"] = report_excel(last["md"], last["computed"], df)
            st.rerun()
        if not (pdf_bytes and docx_bytes):
            st.caption("Missing a download format? Install: `pip install python-docx reportlab`")

# ---------------- Charts ----------------
with tab_charts:
    if not num_cols:
        st.info("No numeric columns detected.")
    else:
        c1, c2, c3, c4 = st.columns(4)
        x = c1.selectbox("Group by (X)", cat_cols + date_cols + [c for c in df.columns if c not in cat_cols + date_cols])
        y = c2.selectbox("Value (Y)", num_cols)
        agg = c3.selectbox("Aggregate", ["sum", "mean", "count", "max", "min"])
        kind = c4.selectbox("Chart", ["Bar", "Line", "Pie"])
        topn = st.slider("Show top N", 5, 50, 15)

        if x == y:
            st.info("Choose a different Group-by column than the value column.")
        else:
            work = df[[x, y]].dropna(subset=[x]).copy()
            if x in date_cols:
                work[x] = work[x].dt.to_period("M").dt.to_timestamp()
            g = work.groupby(x, as_index=False)[y].agg(agg)
            if x in date_cols:
                g = g.sort_values(x)
            else:
                g = g.sort_values(y, ascending=False).head(topn)

            title = f"{agg.title()} of {y} by {x}"
            if kind == "Pie":
                fig = px.pie(g, names=x, values=y, title=title)
            elif kind == "Line" or x in date_cols:
                fig = px.line(g, x=x, y=y, markers=True, title=title)
            else:
                fig = px.bar(g, x=x, y=y, title=title)
            st.plotly_chart(fig, use_container_width=True)

# ---------------- Data ----------------
with tab_data:
    st.dataframe(df.head(500), use_container_width=True)
    st.download_button(
        "Download cleaned data as CSV",
        df.to_csv(index=False).encode("utf-8-sig"),
        "MIS_Data.csv",
        "text/csv",
    )
